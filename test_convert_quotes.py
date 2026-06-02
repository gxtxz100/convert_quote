#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""引号转换单元测试"""

import unittest
from convert_quotes import (
    convert_quotes_in_text,
    QuoteState,
    CN_LEFT_DOUBLE,
    CN_RIGHT_DOUBLE,
    CN_LEFT_SINGLE,
    CN_RIGHT_SINGLE,
    process_paragraph_runs,
)


class TestConvertQuotes(unittest.TestCase):
    def test_simple_double(self):
        out, _ = convert_quotes_in_text('"你好"')
        self.assertEqual(out, f'{CN_LEFT_DOUBLE}你好{CN_RIGHT_DOUBLE}')

    def test_two_phrases_double(self):
        out, _ = convert_quotes_in_text('"甲"和"乙"')
        self.assertEqual(
            out,
            f'{CN_LEFT_DOUBLE}甲{CN_RIGHT_DOUBLE}和{CN_LEFT_DOUBLE}乙{CN_RIGHT_DOUBLE}',
        )

    def test_nested_single_in_double(self):
        out, _ = convert_quotes_in_text('"他说\'你好\'"')
        self.assertEqual(
            out,
            f'{CN_LEFT_DOUBLE}他说{CN_LEFT_SINGLE}你好{CN_RIGHT_SINGLE}{CN_RIGHT_DOUBLE}',
        )

    def test_state_continues_across_chunks(self):
        state = QuoteState()
        part1, state = convert_quotes_in_text('前文"', state)
        part2, state = convert_quotes_in_text('后文"', state)
        self.assertEqual(part1, f'前文{CN_LEFT_DOUBLE}')
        self.assertEqual(part2, f'后文{CN_RIGHT_DOUBLE}')

    def test_word_curly_quotes_normalized(self):
        # 方向错误的弯引号也按顺序纠正
        out, _ = convert_quotes_in_text('\u201d错\u201c')  # 先右后左（错误顺序的输入）
        self.assertEqual(out, f'{CN_LEFT_DOUBLE}错{CN_RIGHT_DOUBLE}')

    def test_straight_quotes_to_chinese(self):
        out, _ = convert_quotes_in_text("'单'")
        self.assertEqual(out, f'{CN_LEFT_SINGLE}单{CN_RIGHT_SINGLE}')

    def test_odd_quote_opens_left(self):
        out, state = convert_quotes_in_text('"未闭合')
        self.assertEqual(out, f'{CN_LEFT_DOUBLE}未闭合')
        self.assertTrue(state.in_double)

    def test_no_change_without_quotes(self):
        text = '普通中文，无引号。'
        out, _ = convert_quotes_in_text(text)
        self.assertEqual(out, text)


class TestDocxParagraph(unittest.TestCase):
    def test_multi_run_paragraph(self):
        from docx import Document

        doc = Document()
        p = doc.add_paragraph()
        p.add_run('甲"').bold = True
        p.add_run('乙')
        p.add_run('"丙')

        state = QuoteState()
        process_paragraph_runs(p, state)

        self.assertEqual(p.text, f'甲{CN_LEFT_DOUBLE}乙{CN_RIGHT_DOUBLE}丙')
        self.assertFalse(state.in_double)

    def test_cross_paragraph_state_in_doc(self):
        """模拟正文连续两段：引号跨段闭合"""
        from docx import Document
        from convert_quotes import _process_block_container

        doc = Document()
        doc.add_paragraph('段一"未完')
        doc.add_paragraph('继续"')

        state = QuoteState()
        _process_block_container(doc, state)

        self.assertEqual(
            doc.paragraphs[0].text,
            f'段一{CN_LEFT_DOUBLE}未完',
        )
        self.assertEqual(
            doc.paragraphs[1].text,
            f'继续{CN_RIGHT_DOUBLE}',
        )


if __name__ == '__main__':
    unittest.main()
